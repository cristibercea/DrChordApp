import logging
from fpdf import FPDF
from docx import Document
from docx.shared import Pt
from odf.text import P
from odf.opendocument import OpenDocumentText
from odf.style import Style, TextProperties, ParagraphProperties

def create_pdf(text_content: str, output_path: str) -> None:
    """
    Creates a PDF file from a text content
    :param text_content: the text data to print in the PDF file
    :param output_path: the path to the output PDF file
    :return: None
    """
    logging.info(f"Creating PDF file...")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Courier", size=9)
    for line in text_content.split('\n'): pdf.cell(200, 4, txt=line, ln=1, align='L')
    pdf.output(output_path)
    logging.info(f"PDF file created at {output_path}.")

def create_docx(text_content: str, output_path: str) -> None:
    """
    Creates a Word file from a text content
    :param text_content: the text data to print in the Word file
    :param output_path: the path to the output Word file
    :return: None
    """
    logging.info(f"Creating Word file...")
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Courier New'
    font.size = Pt(9)
    paragraph_format = style.paragraph_format
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1
    for line in text_content.split('\n'): doc.add_paragraph(line)
    doc.save(output_path)
    logging.info(f"Word file created at {output_path}.")

def create_odt(text_content: str, output_path: str):
    """
    Creates a ODT file from a text content
    :param text_content: the text data to print in the ODT file
    :param output_path: the path to the output ODT file
    :return: None
    """
    logging.info(f"Creating ODT file...")
    doc = OpenDocumentText()
    monostyle = Style(name="MonoTab", family="paragraph")
    monostyle.addElement(TextProperties(fontfamily="Courier New", fontsize="9pt"))
    monostyle.addElement(ParagraphProperties(marginbottom="0cm"))
    doc.styles.addElement(monostyle)
    for line in text_content.split('\n'): doc.text.addElement(P(stylename=monostyle, text=line))
    doc.save(output_path)
    logging.info(f"ODT file created at {output_path}.")
